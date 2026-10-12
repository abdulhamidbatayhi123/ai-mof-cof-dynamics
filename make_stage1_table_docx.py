"""Write the Nature Communications Stage-1 design table as a Word file.

Columns as in the journal's template (nature.com/documents/updated_ncomms-registered-
report-table.docx, read 2026-10-12): Question | Hypotheses | Outcome Measures | Sampling
Plan (N, power analyses) | Analysis Plan | Interpretation given to different outcomes.
The rows restate Table 1 of paper2/stage1.tex, which restates the frozen PREREG_P2
(commit d562f19); power figures from results/p2_power_note.json. A fresh document is
built here; the publisher's file is not copied.

    python make_stage1_table_docx.py   -> paper2/submission/stage1_design_table.docx
"""
import json
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

OUT = "paper2/submission/stage1_design_table.docx"
pw = json.load(open("results/p2_power_note.json", encoding="utf-8"))
p95 = f"{pw['shift_resolvable_at_power_0.95_decade']:.2f}"
p_f2 = f"{pw['power_at_factor_2_shift_log10_2']:.2f}"
d80 = f"{pw['difference_shift_resolvable_at_power_0.80_decade']:.2f}"
d_f2 = f"{pw['difference_power_at_factor_2_shift']:.2f}"

HEAD = ["Question", "Hypotheses", "Outcome Measures", "Sampling Plan (N, power analyses)",
        "Analysis Plan", "Interpretation given to different outcomes"]
ROWS = [
    ["Is the discovery boundary set by the driving force relative to the combined error?",
     "H2a: P(success) plotted against R = δ/ε_eff collapses onto one curve across Da, Pe, "
     "isotherm shape and observation model (O1, O2); the spread of R* across slices is below "
     "half a decade.",
     "Success of the best method per replicate (exact support, correct signs, coefficient "
     "ratio −1 within 10 %); R* (50 % crossing) per slice.",
     "20 replicates per cell; slices are (Pe, isotherm, observation model), O1 and O2, pooling "
     "every σ, ε and Da; cells with R undefined (σ = ε = 0 on O1) excluded. Decided by "
     "thresholds on the spread (0.5 and 1 decade), not by a significance test.",
     "Logistic regression of success on log10 R per slice; R* its 50 % crossing; spread = "
     "max − min of log10 R* with a bootstrap interval; slices without a crossing listed, "
     "never imputed.",
     "Spread < 0.5 decade: “R collapses the boundary”. Spread ≥ 1 decade: “R does not "
     "collapse the boundary; the boundary depends additionally on {named axis}”. Between: "
     "“inconclusive: R narrows the boundary to a spread of {x} decades”."],
    ["Does discovery fail before, with or after the rate constant stops being identifiable?",
     "H2b: compare Da*_disc (best discovery method) with Da*_ident (profile likelihood, true "
     "law known); confirmatory on O1, descriptive on O2 and O3.",
     "Gap log10 Da*_disc − log10 Da*_ident per O1 slice (σ, ε, isotherm).",
     "9 Da levels × 20 replicates for both boundaries. SD of one boundary estimate 0.089 "
     "decade (calibrated on 100 simulated datasets before the freeze): MDE 0.25 decade at "
     f"80 % power, two-sided 5 %; {p95} decade at 95 %. For a gap between two independent "
     f"estimates: {d80} decade at 80 %.",
     "Same logistic estimator for both boundaries; percentile bootstrap (2000 resamples) "
     "within Da levels; the best method re-selected inside every bootstrap draw. O2 and O3 "
     "(3 replicates) reported descriptively.",
     "Gap ≤ −0.25: “discovery fails before identifiability”. |gap| < 0.25: “discovery fails "
     "with identifiability”. Gap ≥ +0.25: “discovery outlives identifiability”. No crossing "
     "in the tested range: no verdict; the range is reported."],
    ["Does an errors-in-variables fit move the boundary when the isotherm is wrong?",
     "H2c: M5 (EIV best-subset) shifts Da*_disc by at least a factor 2 relative to the best "
     "strong-form non-EIV method at ε > 0, on O1.",
     "Shift in log10 Da*_disc, M5 minus the best strong-form non-EIV method, per (isotherm, "
     "ε > 0) at σ = 2 %.",
     "20 replicates; for ε = 2 % a refined grid of 8 extra Da levels within ±0.5 decade of "
     "the best non-EIV boundary. Power to detect the predicted factor 2 (0.30 decade): "
     f"{p_f2} per boundary convention; {d_f2} if the two boundary estimates are independent.",
     "Same estimator and bootstrap; refined grid at ε = 2 %, main grid at ε ∈ {0.5, 5} %; "
     "the shift against the best non-EIV method over both forms reported descriptively.",
     "Shift ≥ log10 2 (0.30): “EIV moves the boundary by {factor}”. MDE ≤ shift < log10 2: “EIV moves the "
     "boundary by {factor}, detectably but by less than the predicted factor 2”. Below the "
     "MDE at every ε > 0: “EIV does not move the boundary detectably (MDE {x})”. Down by at "
     "least the MDE: “EIV moves the boundary DOWN by {factor}”."],
    ["Can a loading-dependent rate at a MOF-like isotherm step be discovered?",
     "H2d (exploratory): no directional prediction.",
     "Discovery of k(q) per cell: support contains q*_meas and q; Spearman ρ < −0.8 across 10 "
     "bins; top-bin rate ≤ 0.6 of bottom-bin rate.",
     "Step isotherm, true law L2, 20 replicates per cell. No power analysis (exploratory).",
     "Rate of discovery per (σ, observation) cell.",
     "Reported as rates; no confirmatory claim rests on H2d."],
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.5))
doc.add_paragraph().add_run(
    "Study design table: When can AI discover adsorption kinetics? A discoverability map "
    "(Registered Report, Stage 1). Every entry restates the frozen pre-registration "
    "(PREREG_P2, commit d562f19).").bold = True
t = doc.add_table(rows=1 + len(ROWS), cols=len(HEAD))
t.style = "Table Grid"
for j, h in enumerate(HEAD):
    t.cell(0, j).text = h
    t.cell(0, j).paragraphs[0].runs[0].bold = True
for i, row in enumerate(ROWS, start=1):
    for j, txt in enumerate(row):
        t.cell(i, j).text = txt
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("wrote", OUT)
