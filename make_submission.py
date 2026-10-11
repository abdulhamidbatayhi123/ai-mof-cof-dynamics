"""Assemble the CMAME submission package and prove it compiles on its own.

Follows the journal's guide for authors as read on 2026-10-11:
  * editable LaTeX source (elsarticle), BibTeX database, every figure as a separate
    file named Figure_1, Figure_2, ... in order of appearance (vector PDF, fonts embedded);
  * highlights as a separate editable file with "highlights" in its name, 3 to 5
    bullets of at most 85 characters each;
  * the graphical abstract as a separate TIFF (>= 1328 x 531 px, readable at 13 x 5 cm);
  * the correction ledger as supplementary material.

The package is rebuilt from the committed sources every time and compiled with
tools/tectonic.exe in a scratch copy; any LaTeX error, undefined reference or missing
file refuses. Run build_paper.py first (it writes numbers.tex and references.bib).

    python make_submission.py   -> submission/ and submission_cmame.zip
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "submission")
N = json.load(open(os.path.join(ROOT, "paper", "numbers.json"), encoding="utf-8"))["values"]

# The highlights, with every number taken from the results (numbers.json).
HIGHLIGHTS = [
    "Seven explanations for surrogate error, tested against pre-registered criteria",
    "A solver checked against closed forms and, unfitted, against a MOF-303 bed",
    "Beyond a few modes the parameter-to-field map binds the error, not the basis",
    f"A coarse solve beats the perceptron surrogates below {N['CoarseBreakEven']} queries",
    "Every number comes from a script; every correction is in a public ledger",
]


def refuse(msg):
    sys.exit("SUBMISSION REFUSED: " + msg)


def main():
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)

    # --- highlights: the guide's limits, checked rather than trusted
    if not 3 <= len(HIGHLIGHTS) <= 5:
        refuse("the guide asks for 3 to 5 highlights")
    for h in HIGHLIGHTS:
        if len(h) > 85:
            refuse(f"highlight over 85 characters ({len(h)}): {h}")
    import docx
    d = docx.Document()
    d.add_heading("Highlights", level=1)
    for h in HIGHLIGHTS:
        d.add_paragraph(h, style="List Bullet")
    d.save(os.path.join(OUT, "highlights.docx"))
    with open(os.path.join(ROOT, "paper", "highlights.txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(HIGHLIGHTS) + "\n")

    # --- manuscript: figures renamed in order of appearance, paths flattened
    tex = open(os.path.join(ROOT, "paper", "manuscript.tex"), encoding="utf-8").read()
    figs = re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", tex)
    mapping = {}
    for i, f in enumerate(figs, 1):
        if f in mapping:
            refuse(f"figure included twice: {f}")
        src = os.path.normpath(os.path.join(ROOT, "paper", f))
        if not os.path.exists(src):
            refuse(f"figure file missing: {src}")
        mapping[f] = f"Figure_{i}.pdf"
        shutil.copy(src, os.path.join(OUT, mapping[f]))
    for f, new in mapping.items():
        tex = tex.replace("{" + f + "}", "{" + new + "}")
    open(os.path.join(OUT, "manuscript.tex"), "w", encoding="utf-8", newline="\n").write(tex)
    # every \input'd file travels with the manuscript (a first version missed one,
    # and only the standalone compile below caught it)
    inputs = [n if n.endswith(".tex") else n + ".tex" for n in re.findall(r"\\input\{([^}]*)\}", tex)]
    for name in inputs + ["references.bib"]:
        src = os.path.join(ROOT, "paper", name)
        if not os.path.exists(src):
            refuse(f"input file missing: {src}")
        shutil.copy(src, OUT)

    # --- graphical abstract and supplementary material
    shutil.copy(os.path.join(ROOT, "figures", "graphical_abstract.tiff"), OUT)
    shutil.copy(os.path.join(ROOT, "paper", "si_ledger.pdf"),
                os.path.join(OUT, "Supplementary_S1_correction_ledger.pdf"))

    # Vitae: biography (editable) and passport photograph, kept in the private outreach/
    # folder and never committed; the package says plainly if either is still missing.
    for name, out_name in (("vitae.docx", "Vitae_Batayhi.docx"), ("photo.jpg", "Vitae_photo_Batayhi.jpg")):
        src = os.path.join(ROOT, "outreach", name)
        if os.path.exists(src):
            shutil.copy(src, os.path.join(OUT, out_name))
        else:
            print(f"NOTE: outreach/{name} not found; the vitae item is incomplete")

    # --- prove the package compiles alone (scratch copy, so no build files land in it)
    tect = os.path.join(ROOT, "tools", "tectonic.exe")
    with tempfile.TemporaryDirectory() as tmp:
        for f in os.listdir(OUT):
            shutil.copy(os.path.join(OUT, f), tmp)
        r = subprocess.run([tect, "--keep-logs", "--keep-intermediates", "manuscript.tex"], cwd=tmp, capture_output=True, text=True)
        log = open(os.path.join(tmp, "manuscript.log"), encoding="utf-8", errors="replace").read() \
            if os.path.exists(os.path.join(tmp, "manuscript.log")) else ""
        bad = [ln for ln in log.splitlines() if ln.startswith("!") or
               ("undefined" in ln.lower() and ("reference" in ln.lower() or "citation" in ln.lower()))]
        if r.returncode != 0 or bad:
            refuse("the package does not compile on its own:\n  " + "\n  ".join(bad[:10] or [r.stderr[-600:]]))
        shutil.copy(os.path.join(tmp, "manuscript.pdf"), os.path.join(OUT, "manuscript_compiled_check.pdf"))
        bbl = os.path.join(tmp, "manuscript.bbl")
        if os.path.exists(bbl):   # for systems that do not run BibTeX
            shutil.copy(bbl, OUT)

    zpath = os.path.join(ROOT, "submission_cmame.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(os.listdir(OUT)):
            z.write(os.path.join(OUT, f), f)
    print("figures:", ", ".join(f"{os.path.basename(k)} -> {v}" for k, v in mapping.items()))
    print("highlights:", ", ".join(str(len(h)) for h in HIGHLIGHTS), "characters")
    print(f"wrote {OUT} ({len(os.listdir(OUT))} files) and {zpath}; compiled standalone without error")


if __name__ == "__main__":
    main()
